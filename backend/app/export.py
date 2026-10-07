"""Rapor dışa aktarma: Excel (openpyxl) ve PDF (reportlab).

Dosyalar kullanıcının İndirilenler\\FinAnaliz klasörüne kaydedilir; arayüz kaydedilen
dosyanın yolunu gösterir ve klasörü açabilir.
"""
import os
from xml.sax.saxutils import escape
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.lineplots import LinePlot
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

EXPORT_DIR = Path.home() / "Downloads" / "FinAnaliz"
DISCLAIMER = "Bu rapor FinAnaliz tarafından yalnızca eğitim ve bilgilendirme amacıyla üretilmiştir; yatırım tavsiyesi değildir."

# Türkçe karakterler için Windows sistem yazı tipi
_FONT, _FONT_BOLD = "Helvetica", "Helvetica-Bold"
for regular, bold in (("arial.ttf", "arialbd.ttf"), ("segoeui.ttf", "segoeuib.ttf"), ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf")):
    fonts = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    if (fonts / regular).exists() and (fonts / bold).exists():
        pdfmetrics.registerFont(TTFont("TR", str(fonts / regular)))
        pdfmetrics.registerFont(TTFont("TR-Bold", str(fonts / bold)))
        _FONT, _FONT_BOLD = "TR", "TR-Bold"
        break

BLUE = colors.HexColor("#2f6fdb")
GREEN = colors.HexColor("#1e8e7e")
RED = colors.HexColor("#d64541")


def _path(name: str, ext: str) -> Path:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    safe = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in name)
    return EXPORT_DIR / f"{safe}_{datetime.now():%Y%m%d_%H%M%S}.{ext}"


def _num(v, digits=2):
    if v is None:
        return "—"
    return f"{v:,.{digits}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _pct(v):
    return "—" if v is None else ("+" if v > 0 else "") + _num(v) + "%"


# ---------------------------------------------------------------- Excel
def excel(title: str, columns: list[tuple[str, str, str]], data: list[dict], name: str, extra_sheets=None) -> Path:
    """columns: (anahtar, başlık, biçim) — biçim: 'text' | 'num' | 'pct' | 'int'."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Oluşturulma: {datetime.now():%d.%m.%Y %H:%M}"])
    ws.append([])
    ws.append([c[1] for c in columns])
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F6FDB")
        cell.alignment = Alignment(horizontal="center")
    for row in data:
        values = []
        for key, _, fmt in columns:
            v = row.get(key)
            if fmt == "pct" and isinstance(v, (int, float)):
                v = v / 100
            values.append(v)
        ws.append(values)
        r = ws.max_row
        for idx, (_, _, fmt) in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=idx)
            if fmt == "pct":
                cell.number_format = "0.00%"
                if isinstance(cell.value, (int, float)):
                    cell.font = Font(color="1E8E7E" if cell.value > 0 else "D64541" if cell.value < 0 else "000000")
            elif fmt == "num":
                cell.number_format = "#,##0.00##"
            elif fmt == "int":
                cell.number_format = "#,##0"
    for idx, (key, head, _) in enumerate(columns, start=1):
        width = max([len(str(head))] + [len(str(r.get(key) or "")) for r in data[:200]]) + 2
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width, 10), 40)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.append([])
    ws.append([DISCLAIMER])
    for sheet_title, cols, sheet_rows in extra_sheets or []:
        s = wb.create_sheet(sheet_title[:31])
        s.append([c[1] for c in cols])
        for cell in s[1]:
            cell.font = Font(bold=True)
        for r in sheet_rows:
            s.append([r.get(c[0]) for c in cols])
        for idx in range(1, len(cols) + 1):
            s.column_dimensions[get_column_letter(idx)].width = 18
    path = _path(name, "xlsx")
    wb.save(path)
    return path


# ---------------------------------------------------------------- PDF
def _styles():
    ss = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("t", parent=ss["Title"], fontName=_FONT_BOLD, fontSize=18, textColor=BLUE, spaceAfter=4),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], fontName=_FONT_BOLD, fontSize=12.5, spaceBefore=10, spaceAfter=6),
        "body": ParagraphStyle("b", parent=ss["Normal"], fontName=_FONT, fontSize=9.5, leading=13),
        "muted": ParagraphStyle("m", parent=ss["Normal"], fontName=_FONT, fontSize=8, textColor=colors.grey, leading=11),
    }


def _table(rows: list[list], col_widths=None, right_cols=None, color_cols=()):
    t = Table(rows, colWidths=col_widths, repeatRows=1)
    style = [
        ("FONTNAME", (0, 0), (-1, -1), _FONT), ("FONTNAME", (0, 0), (-1, 0), _FONT_BOLD),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5), ("BACKGROUND", (0, 0), (-1, 0), BLUE),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f5fa")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d5dbe5")),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for col in (right_cols if right_cols is not None else range(1, len(rows[0]))):
        style.append(("ALIGN", (col, 0), (col, -1), "RIGHT"))
    for col in color_cols:
        for r, row in enumerate(rows[1:], start=1):
            v = str(row[col])
            if v.startswith("+"):
                style.append(("TEXTCOLOR", (col, r), (col, r), GREEN))
            elif v.startswith("-"):
                style.append(("TEXTCOLOR", (col, r), (col, r), RED))
    t.setStyle(TableStyle(style))
    return t


def _line_chart(series: list[list[tuple[float, float]]], line_colors, width=17 * cm, height=6 * cm):
    d = Drawing(width, height)
    lp = LinePlot()
    lp.x, lp.y, lp.width, lp.height = 40, 20, width - 50, height - 30
    lp.data = series
    for i, c in enumerate(line_colors):
        lp.lines[i].strokeColor = c
        lp.lines[i].strokeWidth = 1.4
    lp.xValueAxis.visibleLabels = False
    lp.xValueAxis.visibleTicks = False
    lp.yValueAxis.labels.fontName = _FONT
    lp.yValueAxis.labels.fontSize = 7
    lp.yValueAxis.gridStrokeColor = colors.HexColor("#e3e8ef")
    lp.yValueAxis.visibleGrid = True
    d.add(lp)
    return d


def pdf(name: str, title: str, subtitle: str, blocks: list) -> Path:
    path = _path(name, "pdf")
    st = _styles()
    story = [Paragraph(escape(title), st["title"]), Paragraph(escape(subtitle), st["muted"]), Spacer(1, 8)]
    for kind, content in blocks:
        if kind == "h2":
            story.append(Paragraph(escape(content), st["h2"]))
        elif kind == "p":
            story.append(Paragraph(escape(content), st["body"]))
        elif kind == "legend":
            story.append(Paragraph("   ".join(f'<font color="{c.hexval().replace("0x", "#")}">■</font> {escape(n)}'
                                              for n, c in content), st["body"]))
        elif kind == "table":
            story.append(_table(**content))
        elif kind == "chart":
            story.append(_line_chart(**content))
        story.append(Spacer(1, 4))
    story += [Spacer(1, 12), Paragraph(DISCLAIMER, st["muted"])]
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.5 * cm, title=title, author="FinAnaliz")
    doc.build(story)
    return path


# ---------------------------------------------------------------- Raporlar
def analysis_pdf(symbol: str, name: str, currency: str, candles: list[dict], summary: dict) -> Path:
    st = summary["stats"]
    closes = [(i, c["close"]) for i, c in enumerate(candles)]
    sig_rows = [["Gösterge", "Değer", "Yorum", "Sinyal"]] + [
        [s["name"], _num(s["value"], 4), s["note"], s["signal"].upper()] for s in summary["signals"]]
    stat_rows = [["Ölçüt", "Değer"],
                 ["Son fiyat", f'{_num(st["last"], 4)} {currency}'], ["Günlük değişim", _pct(st["change_1d"])],
                 ["1 aylık değişim", _pct(st["change_1m"])], ["Dönem değişimi (1 yıl)", _pct(st["change_period"])],
                 ["Dönem en yüksek", _num(st["high"], 4)], ["Dönem en düşük", _num(st["low"], 4)],
                 ["Yıllık oynaklık", _pct(st["volatility"]).lstrip("+")], ["ATR (14)", _num(st["atr"], 4)]]
    return pdf(f"Analiz_{symbol}", f"{name} ({symbol}) Teknik Analiz Raporu",
               f"FinAnaliz · {datetime.now():%d.%m.%Y %H:%M} · Son 1 yıllık günlük veri", [
                   ("h2", "Fiyat Grafiği (1 yıl, kapanış)"),
                   ("chart", {"series": [closes], "line_colors": [BLUE]}),
                   ("h2", f"Genel Görünüm: {summary['overall']}"),
                   ("p", "Aşağıdaki tablo, teknik göstergelerin son değerlerine göre otomatik olarak üretilmiş sinyalleri gösterir."),
                   ("table", {"rows": sig_rows, "col_widths": [3.5 * cm, 2.8 * cm, 8 * cm, 2.2 * cm], "right_cols": [1]}),
                   ("h2", "İstatistikler"),
                   ("table", {"rows": stat_rows, "col_widths": [8 * cm, 8.5 * cm], "color_cols": (1,)}),
               ])


def portfolio_pdf(items: list[dict], analytics: dict, username: str) -> Path:
    t = analytics["totals"]
    rows = [["Varlık", "Miktar", "Alış", "Güncel", "Değer", "Kâr/Zarar", "K/Z %"]] + [
        [f'{i["symbol"]}', _num(i["quantity"], 4), _num(i["buy_price"], 4), _num(i["price"], 4),
         f'{_num(i["value"])} {i["currency"]}', _num(i["pnl"]), _pct(i["pnl_pct"])] for i in items]
    blocks = [
        ("h2", "Özet (TL bazında)"),
        ("table", {"rows": [["Ölçüt", "Değer"],
                            ["Toplam değer", f'{_num(t["value_try"])} TL  (≈ {_num(t["value_usd"])} USD)'],
                            ["Toplam maliyet", f'{_num(t["cost_try"])} TL'],
                            ["Kâr / Zarar", f'{"+" if t["pnl_try"] > 0 else ""}{_num(t["pnl_try"])} TL ({_pct(t["pnl_pct"])})'],
                            ["Dönem getirisi (zaman ağırlıklı)", _pct(analytics.get("period_return_pct"))]] +
                   [[f"Kıyas: {k}", _pct(v)] for k, v in analytics.get("benchmark_returns", {}).items()],
                   "col_widths": [8 * cm, 8.5 * cm], "color_cols": (1,)}),
        ("h2", "Pozisyonlar"),
        ("table", {"rows": rows, "color_cols": (5, 6)}),
    ]
    perf = analytics.get("performance", {})
    if perf.get("Portföy"):
        palette = [BLUE, colors.HexColor("#ff9f43"), colors.HexColor("#2ec4b6"), colors.HexColor("#a66cff"), colors.grey]
        names = list(perf.keys())
        blocks += [("h2", "Performans Karşılaştırması (dönem başına göre %)"),
                   ("legend", list(zip(names, palette))),
                   ("chart", {"series": [[(i, p["value"]) for i, p in enumerate(perf[n])] for n in names],
                              "line_colors": palette[:len(names)]})]
    risk = analytics.get("risk", {})
    blocks += [("h2", "Risk Ölçümleri"),
               ("table", {"rows": [["Ölçüt", "Değer"],
                                   ["Yıllık oynaklık", _pct(risk.get("volatility_pct")).lstrip("+")],
                                   ["Maksimum düşüş", _pct(risk.get("max_drawdown_pct"))],
                                   ["Sharpe oranı", _num(risk.get("sharpe"))],
                                   ["Beta (BIST 100)", _num(risk.get("beta_bist100"))],
                                   ["Günlük %95 Riske Maruz Değer (VaR)", f'{_num(risk.get("var95_try"))} TL']],
                          "col_widths": [8 * cm, 8.5 * cm]}),
               ("h2", "Dağılım"),
               ("table", {"rows": [["Piyasa", "Değer (TL)", "Pay"]] + [[a["label"], _num(a["value"]), _num(a["pct"]) + "%"]
                                                                       for a in analytics["allocation"]["market"]]})]
    return pdf("Portfoy_Raporu", f"Portföy Raporu – {username}", f"FinAnaliz · {datetime.now():%d.%m.%Y %H:%M}", blocks)


def open_folder(path: str | None = None) -> None:
    target = Path(path) if path else EXPORT_DIR
    target = target if target.is_dir() else target.parent
    if EXPORT_DIR.resolve() not in [target.resolve(), *target.resolve().parents]:
        raise ValueError("Yalnızca rapor klasörü açılabilir.")
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    os.startfile(str(target))  # noqa: S606 — yerel masaüstü uygulaması
